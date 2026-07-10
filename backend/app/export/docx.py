import os

from docx import Document as Docx
from docx.enum.text import WD_ALIGN_PARAGRAPH

FOOTER = "Δημιουργήθηκε με υποβοήθηση AI — εγκρίθηκε από δικηγόρο. Grafida"


def export_docx(title: str, content: str, out_path: str):
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    d = Docx()
    h = d.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = h.add_run(title)
    run.bold = True
    for para in content.split("\n\n"):
        p = d.add_paragraph(para.strip())
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    f = d.add_paragraph(FOOTER)
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.save(out_path)
